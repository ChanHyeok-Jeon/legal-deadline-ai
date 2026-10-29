"""첨부 파일(HWP, HWPX, DOCX, PDF, TXT)에서 글자만 뽑는다. 파일은 메모리에서만 다루고 저장하지 않는다."""
import io
import re
import struct
import zipfile
import zlib
from xml.etree import ElementTree

MAX_BYTES = 5 * 1024 * 1024
SUPPORTED = (".hwp", ".hwpx", ".docx", ".pdf", ".txt")
IMAGES = (".png", ".jpg", ".jpeg", ".webp", ".gif")


class FileReadError(ValueError):
    pass


# 문단 텍스트 안의 제어 문자: 아래 코드는 뒤따르는 데이터까지 8글자(16바이트)를 차지한다 (HWP 5.0 문서 형식)
_EXTENDED_CTRL = {1, 2, 3, 4, 5, 6, 7, 8, 9, 11, 12, 14, 15, 16, 17, 18, 19, 20, 21, 22, 23}


def _para_text(buf):
    chars, i, n = [], 0, len(buf) // 2
    while i < n:
        c = struct.unpack_from("<H", buf, i * 2)[0]
        if c in _EXTENDED_CTRL:
            i += 8
            continue
        if c in (10, 13):
            chars.append("\n")
        elif c >= 32:
            chars.append(chr(c))
        i += 1
    return "".join(chars)


def _hwp(data):
    """HWP 5.0: OLE 파일 안 BodyText/Section* 스트림의 문단 텍스트 레코드(태그 67)를 읽는다."""
    import olefile
    try:
        ole = olefile.OleFileIO(io.BytesIO(data))
    except OSError:
        raise FileReadError("HWP 파일을 열 수 없습니다. 한글 2002 이전 형식이거나 손상된 파일일 수 있습니다.")
    header = ole.openstream("FileHeader").read()
    if header[36] & 0x02:
        raise FileReadError("암호가 걸린 HWP 파일은 읽을 수 없습니다.")
    compressed = bool(header[36] & 0x01)
    sections = sorted((s for s in ole.listdir() if s[0] == "BodyText"), key=lambda s: int(s[1].replace("Section", "")))
    texts = []
    for s in sections:
        raw = ole.openstream("/".join(s)).read()
        if compressed:
            raw = zlib.decompress(raw, -15)
        i = 0
        while i + 4 <= len(raw):
            head = struct.unpack_from("<I", raw, i)[0]
            tag, size = head & 0x3FF, (head >> 20) & 0xFFF
            i += 4
            if size == 0xFFF:
                size = struct.unpack_from("<I", raw, i)[0]
                i += 4
            if tag == 67:  # HWPTAG_PARA_TEXT
                t = _para_text(raw[i:i + size])
                if t.strip():
                    texts.append(t)
            i += size
    return "\n".join(texts)


def _hwpx(data):
    """HWPX: zip 안 Contents/section*.xml 의 <hp:t> 텍스트."""
    with zipfile.ZipFile(io.BytesIO(data)) as z:
        names = sorted(n for n in z.namelist() if re.match(r"Contents/section\d+\.xml", n))
        out = []
        for n in names:
            root = ElementTree.fromstring(z.read(n))
            for p in root.iter():
                if p.tag.endswith("}p"):
                    line = "".join(t.text or "" for t in p.iter() if t.tag.endswith("}t"))
                    if line.strip():
                        out.append(line)
        return "\n".join(out)


def _docx(data):
    import docx
    d = docx.Document(io.BytesIO(data))
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text for c in row.cells))
    return "\n".join(l for l in lines if l.strip())


def pdf_images(data, limit=3):
    """스캔 PDF용: 앞쪽 페이지에 들어 있는 이미지를 꺼낸다 → [(bytes, format)]"""
    from pypdf import PdfReader
    out = []
    for page in PdfReader(io.BytesIO(data)).pages[:limit]:
        for img in page.images:
            ext = img.name.rsplit(".", 1)[-1].lower()
            fmt_ = {"jpg": "jpeg", "jpeg": "jpeg", "png": "png"}.get(ext)
            if fmt_:
                out.append((img.data, fmt_))
                break
    return out


def _pdf(data):
    from pypdf import PdfReader
    r = PdfReader(io.BytesIO(data))
    if r.is_encrypted:
        raise FileReadError("암호가 걸린 PDF는 읽을 수 없습니다.")
    pages = []
    for p in r.pages[:20]:
        t = p.extract_text(extraction_mode="layout") or ""
        pages.append(re.sub(r"[ \t]{2,}", " ", t))  # 위치 맞춤용 공백 정리
    return "\n".join(pages)


def read_file(filename, data):
    name = filename.lower()
    if len(data) > MAX_BYTES:
        raise FileReadError("파일은 5MB까지 올릴 수 있습니다.")
    if not name.endswith(SUPPORTED):
        raise FileReadError("HWP, HWPX, DOCX, PDF, TXT, 이미지(JPG·PNG) 파일만 올릴 수 있습니다.")
    try:
        if name.endswith(".hwp"):
            text = _hwp(data)
        elif name.endswith(".hwpx"):
            text = _hwpx(data)
        elif name.endswith(".docx"):
            text = _docx(data)
        elif name.endswith(".pdf"):
            text = _pdf(data)
        else:
            text = data.decode("utf-8", errors="ignore")
    except FileReadError:
        raise
    except Exception:
        raise FileReadError("파일을 읽는 중 문제가 생겼습니다. 다른 형식으로 저장해 다시 올려 주세요.")
    text = re.sub(r"[\x00-\x08\x0b-\x1f]", " ", text)  # 일부 PDF는 공백을 제어 문자로 내보낸다
    text = re.sub(r"[ \t]{2,}", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if len(text) < 10:
        msg = "파일에서 글자를 찾지 못했습니다."
        if name.endswith(".pdf"):
            msg += " 스캔한 이미지 PDF는 글자를 읽을 수 없습니다."
        raise FileReadError(msg)
    return text
